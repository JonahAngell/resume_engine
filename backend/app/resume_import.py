"""Parse an uploaded resume into a draft. Nothing is written to the database."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

from pydantic import BaseModel, ConfigDict, Field

PROFILE_ONLY_NOTICE = (
    "This file has your name and summary. Jobs, education, and skills are in "
    "Positions.csv, Education.csv, and Skills.csv from the same LinkedIn download."
)
TEXTLESS_PDF = (
    "This PDF has no selectable text. Upload a text-based PDF or a Word document."
)
TEXTLESS_DOC = "This document has no text. Upload a text-based PDF or a Word document."
BAD_TYPE = "Upload a PDF, Word document (.docx), CSV, or ZIP."
BAD_ZIP = (
    "This ZIP does not contain LinkedIn profile CSV files. Download Profile, "
    "Positions, Education, and Skills from Settings, then Data privacy."
)
BAD_CSV = "This CSV is not a LinkedIn Profile, Positions, Education, or Skills export."
UNREADABLE_PDF = "Could not read this PDF. Upload a text-based PDF or a Word document."
UNREADABLE_DOC = "Could not read this Word document. Save it as .docx and try again."

MONTHS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}
MONTH = (
    "January|February|March|April|June|July|August|September|October|November|"
    "December|Sept|Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec"
)
YEAR = r"(?:19|20)\d{2}"
PART = rf"(?:(?:{MONTH})\.?\s+{YEAR}|{YEAR})"
PRESENT = r"(?:Present|Current|Now)"
DATE_RANGE = re.compile(
    rf"(?P<start>{PART})\s*(?:-|–|—|\bto\b)\s*(?P<end>{PART}|{PRESENT})",
    re.I,
)
YEAR_LINE = re.compile(rf"^(?:(?:{MONTH})\.?\s+)?{YEAR}$", re.I)
EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(?:\+\d{1,3}[\s.-]?)?(?:\(\d{3}\)|\d{3})[\s.-]\d{3}[\s.-]\d{4}")
LINKEDIN_RE = re.compile(r"https?://(?:www\.)?linkedin\.com/\S+", re.I)
BARE_LINKEDIN_RE = re.compile(r"(?:www\.)?linkedin\.com/\S+", re.I)
BULLET_RE = re.compile(r"^\s*(?:[•·●▪◦\-*]|\d+[.)])\s+(\S.*)$")
DEGREE_RE = re.compile(
    r"^(?P<type>(?:ph\.?d\.?|m\.?s\.?|m\.?a\.?|b\.?s\.?|b\.?a\.?|b\.?e\.?|a\.?a\.?|"
    r"associate(?:'s)?(?:\s+of\s+[A-Za-z]+)?|"
    r"bachelor(?:'s)?(?:\s+of\s+[A-Za-z]+)?|"
    r"master(?:'s)?(?:\s+of\s+[A-Za-z]+)?|"
    r"doctorate))"
    r"(?:\s*(?:in|,|–|-)\s+(?P<major>.+))?$",
    re.I,
)
HEADINGS = {
    "summary": "summary",
    "professional summary": "summary",
    "profile": "summary",
    "about": "summary",
    "objective": "summary",
    "experience": "employment",
    "work experience": "employment",
    "professional experience": "employment",
    "employment": "employment",
    "employment history": "employment",
    "work history": "employment",
    "education": "education",
    "education and training": "education",
    "skills": "skills",
    "technical skills": "skills",
    "skills & tools": "skills",
    "tools": "skills",
    "technologies": "skills",
    "core competencies": "skills",
    "projects": "project",
    "selected projects": "project",
    "personal projects": "project",
    "leadership": "leadership",
    "leadership experience": "leadership",
    "leadership and activities": "leadership",
    "volunteer": "leadership",
    "volunteer experience": "leadership",
    "volunteering": "leadership",
    "activities": "leadership",
    "interests": "interests",
    "personal interests": "interests",
    "hobbies": "interests",
}
LINKEDIN_FILES = {"profile.csv", "positions.csv", "education.csv", "skills.csv"}


class ImportProfile(BaseModel):
    model_config = ConfigDict(extra="forbid")
    label: str = "Imported profile"
    name: str = ""
    phone: str = ""
    email: str = ""
    linkedin_label: str = ""
    linkedin_url: str = ""
    included: bool = True
    low_confidence: list[str] = Field(default_factory=list)


class ImportSummary(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str = "Imported summary"
    body: str = Field(default="", max_length=20_000)
    included: bool = True
    low_confidence: list[str] = Field(default_factory=list)


class ImportEntry(BaseModel):
    model_config = ConfigDict(extra="forbid")
    id: str
    section: str = "employment"
    included: bool = True
    source_text: str = Field(default="", max_length=20_000)
    low_confidence: list[str] = Field(default_factory=list)
    organization: str = ""
    title: str = ""
    name: str = ""
    subtitle: str = ""
    location: str = ""
    date_label: str = ""
    start_year: int | None = None
    start_month: int | None = None
    end_year: int | None = None
    end_month: int | None = None
    is_current: bool = False
    bullets: list[str] = Field(default_factory=list, max_length=200)
    institution: str = ""
    college_school: str = ""
    degree_type: str = ""
    degree_major: str = ""
    graduation_year: int | None = None
    graduation_month: int | None = None
    display_date: str = ""
    tools: list[str] = Field(default_factory=list, max_length=500)
    body: str = Field(default="", max_length=20_000)


class ImportDraft(BaseModel):
    model_config = ConfigDict(extra="forbid")
    profile: ImportProfile = Field(default_factory=ImportProfile)
    summary: ImportSummary = Field(default_factory=ImportSummary)
    entries: list[ImportEntry] = Field(default_factory=list, max_length=1000)
    leftover: str = Field(default="", max_length=200_000)
    notices: list[str] = Field(default_factory=list, max_length=100)


@dataclass
class TextLine:
    text: str
    bold: bool = False
    size: float = 0
    bullet: bool = False


@dataclass
class _Date:
    label: str = ""
    start_year: int | None = None
    start_month: int | None = None
    end_year: int | None = None
    end_month: int | None = None
    is_current: bool = False
    found: bool = False


MAX_ARCHIVE_ENTRIES = 50
MAX_ARCHIVE_MEMBER_BYTES = 5 * 1024 * 1024
MAX_ARCHIVE_TOTAL_BYTES = 20 * 1024 * 1024
MAX_CSV_ROWS = 5000
TOO_LARGE = "This file expands to more data than the importer allows."


def _check_archive(archive: zipfile.ZipFile) -> None:
    """Reject zip bombs by trusting the declared sizes before anything is decompressed."""
    infos = archive.infolist()
    if len(infos) > MAX_ARCHIVE_ENTRIES:
        raise ValueError(TOO_LARGE)
    total = 0
    for info in infos:
        total += info.file_size
        if info.file_size > MAX_ARCHIVE_MEMBER_BYTES or total > MAX_ARCHIVE_TOTAL_BYTES:
            raise ValueError(TOO_LARGE)


def _read_member(archive: zipfile.ZipFile, info: zipfile.ZipInfo) -> bytes:
    """Read one member, stopping at the cap even if the header understates its size."""
    with archive.open(info) as handle:
        data = handle.read(MAX_ARCHIVE_MEMBER_BYTES + 1)
    if len(data) > MAX_ARCHIVE_MEMBER_BYTES:
        raise ValueError(TOO_LARGE)
    return data


def parse_upload(filename: str, data: bytes) -> ImportDraft:
    if not data:
        raise ValueError("Choose a file to import.")
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        return parse_text_resume(extract_pdf(data))
    if ext == ".docx":
        return parse_text_resume(extract_docx(data))
    if ext == ".csv":
        return parse_linkedin_csv(data)
    if ext == ".zip":
        return parse_linkedin_zip(data)
    raise ValueError(BAD_TYPE)


def parse_text_resume(lines: list[TextLine]) -> ImportDraft:
    useful = [line for line in lines if line.text.strip()]
    if not useful:
        raise ValueError(TEXTLESS_DOC)
    preamble, sections, unknown = _split_sections(useful)
    draft = ImportDraft()
    summary_parts: list[str] = []
    leftover: list[str] = []
    for kind, block in sections:
        if kind == "summary":
            summary_parts.extend(_plain(block))
        elif kind == "skills":
            draft.entries.extend(_skills(block))
        elif kind == "interests":
            draft.entries.extend(_interests(block))
        elif kind in {"employment", "leadership", "project", "education"}:
            draft.entries.extend(_dated_entries(block, kind))
        else:
            leftover.extend(_plain(block))
    leftover.extend(unknown)
    contact, extra = _contact(preamble)
    draft.profile = contact
    if summary_parts:
        draft.summary.body = " ".join(summary_parts).strip()
    elif extra:
        draft.summary.body = " ".join(extra).strip()
        extra = []
    if extra:
        leftover = extra + leftover
    if not draft.summary.body:
        draft.summary.included = False
    draft.leftover = "\n".join(part for part in leftover if part).strip()
    return draft


def extract_pdf(data: bytes) -> list[TextLine]:
    try:
        from pdfminer.high_level import extract_pages
        from pdfminer.layout import LTChar, LTTextContainer, LTTextLine
    except ImportError as exc:  # pragma: no cover
        raise ValueError(UNREADABLE_PDF) from exc
    lines: list[TextLine] = []
    try:
        for page in extract_pages(io.BytesIO(data)):
            for element in page:
                if not isinstance(element, LTTextContainer):
                    continue
                for line in element:
                    if not isinstance(line, LTTextLine):
                        continue
                    text = line.get_text().strip()
                    if not text:
                        continue
                    size = 0.0
                    fonts: list[str] = []
                    for obj in line:
                        if isinstance(obj, LTChar):
                            size = max(size, float(obj.size or 0))
                            fonts.append(obj.fontname or "")
                    bold = any("bold" in font.lower() for font in fonts)
                    lines.append(TextLine(text, bold=bold, size=size, bullet=_is_bullet(text)))
    except Exception as exc:
        raise ValueError(UNREADABLE_PDF) from exc
    if not lines:
        raise ValueError(TEXTLESS_PDF)
    return lines


def extract_docx(data: bytes) -> list[TextLine]:
    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError as exc:  # pragma: no cover
        raise ValueError(UNREADABLE_DOC) from exc
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            _check_archive(package)
    except zipfile.BadZipFile as exc:
        raise ValueError(UNREADABLE_DOC) from exc
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(UNREADABLE_DOC) from exc
    lines: list[TextLine] = []

    def add_paragraph(paragraph) -> None:
        text = paragraph.text.strip()
        if not text:
            return
        style = paragraph.style.name if paragraph.style is not None else ""
        bold = bool(paragraph.runs) and all(run.bold for run in paragraph.runs if run.text.strip())
        bullet = style.lower().startswith("list") or _is_bullet(text) or _numbering(paragraph, qn)
        lines.append(TextLine(text, bold=bold or style.lower().startswith("heading"), bullet=bullet))

    for paragraph in document.paragraphs:
        add_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    add_paragraph(paragraph)
    if not lines:
        raise ValueError(TEXTLESS_DOC)
    return lines


def parse_linkedin_zip(data: bytes) -> ImportDraft:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ValueError(BAD_ZIP) from exc
    _check_archive(archive)
    files: dict[str, bytes] = {}
    for info in archive.infolist():
        if info.is_dir():
            continue
        name = Path(info.filename).name.lower()
        if name in LINKEDIN_FILES and name not in files:
            files[name] = _read_member(archive, info)
    if not files:
        raise ValueError(BAD_ZIP)
    return _linkedin_draft(files)


def parse_linkedin_csv(data: bytes) -> ImportDraft:
    rows, headers = _csv_rows(data)
    kind = _csv_kind(headers)
    if not kind:
        raise ValueError(BAD_CSV)
    return _linkedin_draft({f"{kind}.csv": data})


def _linkedin_draft(files: dict[str, bytes]) -> ImportDraft:
    draft = ImportDraft()
    present = set(files)
    if "profile.csv" in files:
        rows, _ = _csv_rows(files["profile.csv"])
        if rows:
            _apply_profile(draft, rows[0])
    if "positions.csv" in files:
        rows, _ = _csv_rows(files["positions.csv"])
        for row in rows:
            entry = _position(row)
            if entry:
                draft.entries.append(entry)
    if "education.csv" in files:
        rows, _ = _csv_rows(files["education.csv"])
        for row in rows:
            entry = _education_row(row)
            if entry:
                draft.entries.append(entry)
    if "skills.csv" in files:
        rows, _ = _csv_rows(files["skills.csv"])
        tools = [_col(row, "name", "skill") for row in rows]
        tools = [tool for tool in tools if tool]
        if tools:
            draft.entries.append(_entry("skills", name="Skills", tools=tools, source_text="\n".join(tools)))
    missing = [name for name in ("positions.csv", "education.csv", "skills.csv") if name not in present]
    if missing == ["positions.csv", "education.csv", "skills.csv"]:
        draft.notices.append(PROFILE_ONLY_NOTICE)
    elif missing:
        joined = ", ".join(name for name in missing)
        draft.notices.append(f"Missing from this download: {joined}.")
    if not draft.summary.body:
        draft.summary.included = False
    return draft


def _apply_profile(draft: ImportDraft, row: dict[str, str]) -> None:
    first = _col(row, "first name")
    last = _col(row, "last name")
    name = " ".join(part for part in (first, last) if part).strip() or _col(row, "name")
    draft.profile.name = name
    draft.profile.email = _col(row, "email", "email address")
    draft.profile.phone = _col(row, "phone", "phone number")
    websites = _col(row, "websites", "website")
    linkedin = LINKEDIN_RE.search(websites) or BARE_LINKEDIN_RE.search(websites)
    if linkedin:
        draft.profile.linkedin_url = _http(linkedin.group(0))
        draft.profile.linkedin_label = "LinkedIn"
    headline = _col(row, "headline")
    summary = _col(row, "summary", "about")
    if headline:
        draft.summary.name = headline[:120]
    if summary:
        draft.summary.body = summary
    elif headline:
        draft.summary.body = headline
    if not draft.profile.name:
        draft.profile.low_confidence.append("name")


def _position(row: dict[str, str]) -> ImportEntry | None:
    organization = _col(row, "company name", "company", "organization")
    title = _col(row, "title")
    if not organization and not title:
        return None
    location = _col(row, "location")
    description = _col(row, "description", "job description")
    started = _col(row, "started on", "start date")
    finished = _col(row, "finished on", "end date")
    parsed = _from_endpoints(started, finished)
    low: list[str] = []
    if not organization:
        low.append("organization")
    if not title:
        low.append("title")
    return _entry(
        "employment",
        organization=organization,
        title=title,
        location=location,
        date_label=parsed.label,
        start_year=parsed.start_year,
        start_month=parsed.start_month,
        end_year=parsed.end_year,
        end_month=parsed.end_month,
        is_current=parsed.is_current,
        bullets=_description_bullets(description),
        source_text="\n".join(part for part in (organization, title, location, parsed.label, description) if part),
        low_confidence=low,
    )


def _education_row(row: dict[str, str]) -> ImportEntry | None:
    institution = _col(row, "school name", "institution", "school")
    degree_name = _col(row, "degree name", "degree")
    if not institution and not degree_name:
        return None
    degree_type, degree_major = _split_degree(degree_name)
    if degree_name and not degree_type:
        degree_type = degree_name
    started = _col(row, "start date", "started on")
    finished = _col(row, "end date", "finished on")
    parsed = _from_endpoints(started, finished)
    graduation_year = parsed.end_year or parsed.start_year
    graduation_month = parsed.end_month if parsed.end_year else parsed.start_month
    notes = _col(row, "notes", "activities")
    low = ["institution"] if not institution else []
    return _entry(
        "education",
        institution=institution,
        degree_type=degree_type,
        degree_major=degree_major,
        graduation_year=graduation_year,
        graduation_month=graduation_month,
        display_date=parsed.label,
        bullets=_description_bullets(notes),
        source_text="\n".join(part for part in (institution, degree_name, parsed.label, notes) if part),
        low_confidence=low,
    )


def _description_bullets(text: str) -> list[str]:
    bullets: list[str] = []
    for line in re.split(r"[\r\n]+", text or ""):
        cleaned = line.strip()
        if not cleaned:
            continue
        bullets.append(_strip_bullet(cleaned) if _is_bullet(cleaned) else cleaned)
    return bullets


def _csv_rows(data: bytes) -> tuple[list[dict[str, str]], list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    headers = [header or "" for header in (reader.fieldnames or [])]
    rows = []
    for row in reader:
        if len(rows) >= MAX_CSV_ROWS:
            raise ValueError(TOO_LARGE)
        rows.append({(key or "").strip(): (value or "").strip() for key, value in row.items()})
    return rows, headers


def _csv_kind(headers: list[str]) -> str | None:
    names = {header.strip().lower() for header in headers}
    if "company name" in names or ("title" in names and ("started on" in names or "start date" in names)):
        return "positions"
    if "school name" in names or "degree name" in names:
        return "education"
    if "first name" in names or "headline" in names or "summary" in names:
        return "profile"
    if names == {"name"} or "skill" in names:
        return "skills"
    return None


def _col(row: dict[str, str], *names: str) -> str:
    lowered = {key.lower(): value for key, value in row.items()}
    for name in names:
        value = lowered.get(name.lower(), "").strip()
        if value:
            return value
    return ""


def _split_sections(lines: list[TextLine]) -> tuple[list[TextLine], list[tuple[str, list[TextLine]]], list[str]]:
    preamble: list[TextLine] = []
    sections: list[tuple[str, list[TextLine]]] = []
    mode = "preamble"
    kind: str | None = None
    block: list[TextLine] = []

    def flush() -> None:
        nonlocal block
        if mode == "section" and kind:
            sections.append((kind, block))
        elif mode == "preamble":
            preamble.extend(block)
        block = []

    for line in lines:
        heading = None if line.bullet else _heading_kind(line)
        if heading:
            flush()
            mode = "section"
            kind = heading
            continue
        block.append(line)
    flush()
    return preamble, sections, []


def _heading_kind(line: TextLine) -> str | None:
    key = re.sub(r"\s+", " ", line.text.strip().lower()).rstrip(":").strip()
    kind = HEADINGS.get(key)
    if not kind:
        return None
    if line.size and line.size < 11 and not line.bold and len(key.split()) > 2:
        return None
    return kind


def _dated_entries(lines: list[TextLine], section: str) -> list[ImportEntry]:
    groups = _group_by_date(lines)
    if not groups:
        return []
    return [_role(group, section) if section != "education" else _education_block(group) for group in groups]


def _group_by_date(lines: list[TextLine]) -> list[list[TextLine]]:
    indexes = [i for i, line in enumerate(lines) if not line.bullet and _is_date_line(line.text)]
    if not indexes:
        return [lines] if lines else []
    groups: list[list[TextLine]] = []
    for number, date_at in enumerate(indexes):
        previous = indexes[number - 1] + 1 if number else 0
        start = date_at
        while start > previous and not lines[start - 1].bullet:
            start -= 1
        if number + 1 < len(indexes):
            nxt = indexes[number + 1]
            end = nxt
            while end > date_at + 1 and not lines[end - 1].bullet:
                end -= 1
        else:
            end = len(lines)
        group = lines[start:end]
        if group:
            groups.append(group)
    return groups


def _role(lines: list[TextLine], section: str) -> ImportEntry:
    bullets, texts = _bullets_and_text(lines)
    parsed, texts = _take_date(texts)
    location, texts = _take_location(texts)
    headers, extras = _headers(texts)
    bullets = extras + bullets
    low: list[str] = []
    if not parsed.found:
        low.append("start_year")
    if section == "project":
        name = headers[0] if headers else ""
        subtitle = headers[1] if len(headers) > 1 else ""
        if not name:
            low.append("name")
        return _entry(
            section,
            name=name,
            subtitle=subtitle,
            location=location,
            bullets=bullets,
            source_text=_source(lines),
            low_confidence=low,
            **_date_kwargs(parsed),
        )
    organization = headers[0] if headers else ""
    title = headers[1] if len(headers) > 1 else ""
    if not organization:
        low.append("organization")
    if not title:
        low.append("title")
    return _entry(
        section,
        organization=organization,
        title=title,
        location=location,
        bullets=bullets,
        source_text=_source(lines),
        low_confidence=low,
        **_date_kwargs(parsed),
    )


def _education_block(lines: list[TextLine]) -> ImportEntry:
    bullets, texts = _bullets_and_text(lines)
    parsed, texts = _take_date(texts)
    location, texts = _take_location(texts)
    degree_type = ""
    degree_major = ""
    rest: list[str] = []
    for text in texts:
        found = _split_degree(text)
        if found[0] and not degree_type:
            degree_type, degree_major = found
        else:
            rest.append(text)
    institution = rest[0] if rest else ""
    college = rest[1] if len(rest) > 1 else ""
    year = parsed.end_year or parsed.start_year
    month = parsed.end_month if parsed.end_year else parsed.start_month
    low = []
    if not institution:
        low.append("institution")
    if not parsed.found:
        low.append("graduation_year")
    return _entry(
        "education",
        institution=institution,
        college_school=college,
        location=location,
        degree_type=degree_type,
        degree_major=degree_major,
        graduation_year=year,
        graduation_month=month,
        display_date=parsed.label,
        bullets=bullets,
        source_text=_source(lines),
        low_confidence=low,
    )


def _skills(lines: list[TextLine]) -> list[ImportEntry]:
    entries: list[ImportEntry] = []
    loose: list[str] = []

    def flush() -> None:
        if loose:
            entries.append(_entry("skills", name="Skills", tools=list(loose), source_text="\n".join(loose)))
            loose.clear()

    for line in lines:
        text = _strip_bullet(line.text) if line.bullet or _is_bullet(line.text) else line.text.strip()
        if text and ":" in text and not line.bullet:
            flush()
            label, rest = text.split(":", 1)
            tools = _split_tools(rest)
            if tools:
                entries.append(
                    _entry("skills", name=(label.strip() or "Skills")[:120], tools=tools, source_text=text)
                )
        else:
            loose.extend(_split_tools(text) if "," in text else ([text] if text else []))
    flush()
    return entries


def _interests(lines: list[TextLine]) -> list[ImportEntry]:
    entries: list[ImportEntry] = []
    for line in lines:
        text = _strip_bullet(line.text) if _is_bullet(line.text) or line.bullet else line.text.strip()
        if not text:
            continue
        parts = [part.strip() for part in text.split(",")] if text.count(",") and len(text.split()) <= 12 else [text]
        for part in parts:
            if part:
                entries.append(_entry("interests", body=part, source_text=part))
    return entries


def _contact(lines: list[TextLine]) -> tuple[ImportProfile, list[str]]:
    blob = "\n".join(line.text for line in lines)
    email = EMAIL_RE.search(blob)
    phone = PHONE_RE.search(blob)
    linkedin = LINKEDIN_RE.search(blob) or BARE_LINKEDIN_RE.search(blob)
    names: list[str] = []
    for line in lines:
        cleaned = _strip_contact(line.text)
        if cleaned:
            names.append(cleaned)
    profile = ImportProfile(
        name=names[0] if names else "",
        email=email.group(0) if email else "",
        phone=phone.group(0) if phone else "",
        linkedin_url=_http(linkedin.group(0)) if linkedin else "",
        linkedin_label="LinkedIn" if linkedin else "",
    )
    if not profile.name:
        profile.low_confidence.append("name")
    return profile, names[1:]


def _strip_contact(text: str) -> str:
    text = EMAIL_RE.sub(" ", text)
    text = PHONE_RE.sub(" ", text)
    text = LINKEDIN_RE.sub(" ", text)
    text = BARE_LINKEDIN_RE.sub(" ", text)
    text = re.sub(r"[|•·]+", " ", text)
    return re.sub(r"\s+", " ", text).strip(" -")


def _http(url: str) -> str:
    url = url.strip().rstrip(".,)")
    if url.lower().startswith(("http://", "https://")):
        return url
    return f"https://{url}"


def _bullets_and_text(lines: list[TextLine]) -> tuple[list[str], list[str]]:
    bullets: list[str] = []
    texts: list[str] = []
    for line in lines:
        if line.bullet or _is_bullet(line.text):
            bullets.append(_strip_bullet(line.text))
        elif line.text.strip():
            texts.append(line.text.strip())
    return bullets, texts


def _take_date(texts: list[str]) -> tuple[_Date, list[str]]:
    parsed = _Date()
    kept: list[str] = []
    for text in texts:
        if not parsed.found:
            match = DATE_RANGE.search(text)
            if match:
                parsed = _from_range(match.group(0), match.group("start"), match.group("end"))
                rest = (text[: match.start()] + " " + text[match.end() :]).strip(" |-–—·")
                rest = _drop_duration(rest)
                if rest:
                    kept.append(rest)
                continue
            if _is_date_line(text) and len(text) <= 32:
                parsed = _from_endpoints(text, "", open_ended_is_current=False)
                continue
        kept.append(text)
    return parsed, kept


def _from_range(label: str, start: str, end: str) -> _Date:
    start_year, start_month, _ = _parse_point(start)
    if re.fullmatch(PRESENT, end.strip(), re.I):
        return _Date(label.strip(), start_year, start_month, None, None, True, True)
    end_year, end_month, _ = _parse_point(end)
    return _Date(label.strip(), start_year, start_month, end_year, end_month, False, True)


def _from_endpoints(start: str, end: str, open_ended_is_current: bool = True) -> _Date:
    start_year, start_month, _ = _parse_point(start) if start else (None, None, False)
    finished = (end or "").strip()
    present = bool(finished) and bool(re.fullmatch(PRESENT, finished, re.I))
    current = present or (open_ended_is_current and bool(start) and not finished)
    if current:
        end_year = end_month = None
        label_end = "Present" if start else ""
    else:
        end_year, end_month, _ = _parse_point(finished) if finished else (None, None, False)
        label_end = finished
    label = " - ".join(part for part in (start.strip(), label_end) if part)
    found = bool(start_year or end_year or current)
    return _Date(label, start_year, start_month, end_year, end_month, current, found)


def _parse_point(text: str) -> tuple[int | None, int | None, bool]:
    raw = text.strip().rstrip(".")
    if re.fullmatch(PRESENT, raw, re.I):
        return None, None, True
    iso = re.fullmatch(r"(\d{4})-(\d{2})(?:-\d{2})?", raw)
    if iso:
        month = int(iso.group(2))
        if 1 <= month <= 12:
            return int(iso.group(1)), month, False
    named = re.fullmatch(rf"({MONTH})\.?\s+({YEAR})", raw, re.I)
    if named:
        return int(named.group(2)), MONTHS[named.group(1).lower().rstrip(".")], False
    year = re.fullmatch(YEAR, raw)
    if year:
        return int(year.group(0)), None, False
    return None, None, False


def _take_location(texts: list[str]) -> tuple[str, list[str]]:
    for index, text in enumerate(texts):
        if _looks_like_location(text):
            return text, texts[:index] + texts[index + 1 :]
    return "", texts


def _headers(texts: list[str]) -> tuple[list[str], list[str]]:
    headers: list[str] = []
    extras: list[str] = []
    for text in texts:
        cleaned = _drop_duration(text)
        if not cleaned:
            continue
        if len(headers) >= 2 or (headers and len(cleaned) > 80):
            extras.append(cleaned)
        else:
            headers.append(cleaned)
    return headers, extras


def _looks_like_location(text: str) -> bool:
    if len(text) > 80 or any(character.isdigit() for character in text):
        return False
    return "," in text and len(text.split()) <= 10


def _split_degree(text: str) -> tuple[str, str]:
    match = DEGREE_RE.match(text.strip())
    if not match:
        return "", ""
    return match.group("type").strip(), (match.group("major") or "").strip(" .")


def _split_tools(text: str) -> list[str]:
    parts = re.split(r"[,;|•·]", text)
    return [part.strip(" -*") for part in parts if part.strip(" -*")]


def _is_bullet(text: str) -> bool:
    return bool(BULLET_RE.match(text))


def _strip_bullet(text: str) -> str:
    match = BULLET_RE.match(text)
    return match.group(1).strip() if match else text.strip()


def _is_date_line(text: str) -> bool:
    if DATE_RANGE.search(text):
        return True
    return bool(YEAR_LINE.fullmatch(text.strip()))


def _drop_duration(text: str) -> str:
    cleaned = text.strip(" |-–—·")
    if re.fullmatch(
        r"\d+\s*(?:yrs?|mos?|years?|months?)(?:\s+\d+\s*(?:yrs?|mos?|years?|months?))?",
        cleaned,
        re.I,
    ):
        return ""
    return cleaned


def _plain(lines: list[TextLine]) -> list[str]:
    return [line.text.strip() for line in lines if line.text.strip()]


def _source(lines: list[TextLine]) -> str:
    return "\n".join(line.text.strip() for line in lines if line.text.strip())


def _date_kwargs(parsed: _Date) -> dict:
    return {
        "date_label": parsed.label,
        "start_year": parsed.start_year,
        "start_month": parsed.start_month,
        "end_year": parsed.end_year,
        "end_month": parsed.end_month,
        "is_current": parsed.is_current,
    }


def _entry(section: str, **kwargs) -> ImportEntry:
    low = kwargs.pop("low_confidence", [])
    return ImportEntry(id=uuid4().hex, section=section, low_confidence=low, **kwargs)


def _numbering(paragraph, qn) -> bool:
    p_pr = paragraph._p.pPr
    return bool(p_pr is not None and p_pr.find(qn("w:numPr")) is not None)
